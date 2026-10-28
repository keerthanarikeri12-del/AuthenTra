from pathlib import Path
import pdfplumber
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path('tmp/report'); P.mkdir(exist_ok=True,parents=True)
Path('output/report').mkdir(exist_ok=True,parents=True)
src=pdfplumber.open(r'C:\Users\keert\Downloads\record template.pdf')
for i in range(4): src.pages[i].to_image(resolution=200).save(str(P/f'front{i}.png'))
bg=src.pages[5].to_image(resolution=200).original.convert('RGB'); dr=ImageDraw.Draw(bg); scale=200/72; dr.rectangle((27*scale,27*scale,568*scale,815*scale),fill='white'); bg.save(P/'border.png')
doc=Document(); sec=doc.sections[0]
sec.page_width=Pt(595.2); sec.page_height=Pt(841.92); sec.top_margin=Pt(32); sec.bottom_margin=Pt(45); sec.left_margin=Pt(34.1); sec.right_margin=Pt(34.1); sec.header_distance=Pt(0); sec.footer_distance=Pt(31)
n=doc.styles['Normal']; n.font.name='Times New Roman'; n.font.size=Pt(12); n.paragraph_format.line_spacing=Pt(20.64); n.paragraph_format.space_after=Pt(10); n.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(17 if name=='Heading 1' else 14); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(6); s.paragraph_format.space_after=Pt(8)
def anchor(p,path):
 sh=p.add_run().add_picture(str(path),width=Pt(595.2),height=Pt(841.92)); il=sh._inline; a=OxmlElement('wp:anchor')
 for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items(): a.set(k,v)
 e=OxmlElement('wp:simplePos'); e.set('x','0'); e.set('y','0'); a.append(e)
 for axis in ['H','V']:
  e=OxmlElement('wp:position'+axis); e.set('relativeFrom','page'); off=OxmlElement('wp:posOffset'); off.text='0'; e.append(off); a.append(e)
 for child in list(il): a.append(child)
 a.insert(4,OxmlElement('wp:wrapNone')); il.getparent().replace(il,a)
for i in range(4):
 anchor(doc.add_paragraph(),P/f'front{i}.png')
 if i<3: doc.add_page_break()
sec=doc.add_section(WD_SECTION_START.NEW_PAGE); sec.header.is_linked_to_previous=False; sec.footer.is_linked_to_previous=False; anchor(sec.header.paragraphs[0],P/'border.png')
p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); r.font.size=Pt(11); r.font.color.rgb=RGBColor(0,80,255); f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),'begin'); r._r.append(f); r=p.add_run(); x=OxmlElement('w:instrText'); x.text=' PAGE '; r._r.append(x); r=p.add_run(); f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),'end'); r._r.append(f)
font=ImageFont.truetype('C:/Windows/Fonts/times.ttf',27)
def chart(kind):
 im=Image.new('RGB',(1400,1560),'white'); d=ImageDraw.Draw(im)
 def box(x,y,w,h,txt,diamond=False):
  if diamond: d.polygon([(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)],outline='black',width=3)
  else: d.rectangle((x,y,x+w,y+h),outline='black',width=3)
  lines=txt.split('\n')
  for j,t in enumerate(lines): d.text((x+w/2,y+h/2+(j-(len(lines)-1)/2)*33),t,font=font,fill='black',anchor='mm')
 def arrow(x,y,x2,y2,label=''):
  d.line((x,y,x2,y2),fill='black',width=3)
  if y2>y: d.polygon([(x2-8,y2-12),(x2+8,y2-12),(x2,y2)],fill='black')
  else: d.polygon([(x2-12,y2-8),(x2-12,y2+8),(x2,y2)],fill='black')
  if label:d.text(((x+x2)/2-8,(y+y2)/2-35),label,font=font,fill='black')
 if kind=='ARCH':
  box(100,30,1200,120,'Verifier | Institution | System administrator\nReact / Vite interface'); arrow(700,150,700,210)
  box(100,210,1200,120,'FastAPI\nAuthentication | Permissions | File validation')
  for x in [230,700,1170]: arrow(700,330,x,450)
  box(30,450,400,230,'Blockchain adapter\nSHA-256 of original bytes\nSolidity registry\nGanache local chain')
  box(500,450,400,230,'SQLite off-chain store\nAccounts and institutions\nConfirmed metadata\nHistory and reviews')
  box(970,450,400,230,'Advisory analysis\nOCR and Random Forest\nField and visual checks\nTransformation analysis')
  for x in [230,700,1170]:arrow(x,680,700,850)
  box(100,850,1200,140,'Response and review routing\nAuthoritative outcome + separate advisory evidence')
  d.text((700,1080),'Only registry checks and exact hashes determine verification.',font=font,fill='black',anchor='mm')
  im=im.crop((0,0,1400,1150))
 else:
  box(200,10,650,75,'Sign in -> Upload document'); arrow(525,85,525,120)
  checks=[('Valid file?', 'INVALID FILE'),('Registry reachable?', 'UNAVAILABLE'),('Record found?', 'UNREGISTERED /\\nINCONCLUSIVE'),('Active record?', 'REVOKED'),('Issuer approved?', 'ISSUER NOT AUTHORIZED'),('Exact hash match?', 'FILE MISMATCH')]
  y=120
  for title,out in checks:
   box(200,y,650,130,title,True); arrow(850,y+65,920,y+65,'No'); box(920,y+25,460,80,out); arrow(525,y+130,525,y+180,'Yes'); y+=180
  box(200,y,650,75,'VERIFIED'); arrow(525,y+75,525,y+125)
  box(80,y+125,1280,110,'Display primary result and separate advisory findings\nRecord history -> Review when needed -> End')
  im=im.crop((0,0,1400,y+250))
 im.save(P/(kind+'.png')); p=doc.add_paragraph(); p.paragraph_format.line_spacing=1; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(P/(kind+'.png')),width=Pt(510))
def table():
 rows=[['Title','Author(s)','Year','Objective','Advantages','Disadvantages','Literature Survey Summary'],['Blockcerts Guide [1]','Blockcerts community','n.d.','Verify blockchain credentials','Open tools','Requires issuer trust','Informs digest checks; custom registry differs.'],['VC Data Model v2.0 [2]','W3C','2025','Represent verifiable claims','Shared data model','Needs proofs and exchange protocols','Guides roles and privacy; full compliance is future work.'],['Tesseract OCR [3]','Ray Smith','2007','Describe OCR','Extracts text','Quality-dependent errors','Text extraction is not authentication.'],['Random Forests [4]','Leo Breiman','2001','Ensemble classification','Multiple decision trees','Depends on training data','Supports experimental advisory classification.']]
 t=doc.add_table(rows=0,cols=7); t.autofit=False
 for row in rows:
  cells=t.add_row().cells
  for i,txt in enumerate(row):
   p=cells[i].paragraphs[0]; p.paragraph_format.line_spacing=Pt(15); p.paragraph_format.space_after=Pt(8); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run(txt); r.font.size=Pt(10); r.bold=(row==rows[0])
 b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'000000'); b.append(e)
 t._tbl.tblPr.append(b)
pages=Path('tmp/report_content.txt').read_text(encoding='utf-8').split('===')
nums=[6,6,6,6,6,7,8,11,12,13,15,17,18,18,19,19,20,21]
for idx,content in enumerate(pages):
 if idx:doc.add_page_break()
 for j,line in enumerate(content.strip().splitlines()):
  if idx==0:
   if j==0:doc.add_paragraph(line,'Heading 2');continue
   p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.tab_stops.add_tab_stop(Pt(510),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS); p.add_run(line+'\t'+str(nums[j-1]));continue
  if line=='@LITERATURE':table()
  elif line.startswith('@'):chart(line[1:])
  elif line.startswith('##'):doc.add_paragraph(line[2:],'Heading 2')
  elif line.startswith('#'):
   p=doc.add_paragraph(line[1:],'Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  elif line:doc.add_paragraph(line)
doc.core_properties.title='AuthenTra - AI assisted digital certificate authenticity verification using blockchain'
doc.save('output/report/AuthenTra_Mini_Project_Report.docx')
print('Saved',len(pages)+4,'planned pages')

